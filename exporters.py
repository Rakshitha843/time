"""Excel and PDF export of a timetable, one grid per student group, teacher or room."""
from __future__ import annotations

import io
from collections import defaultdict

from models import Config, is_batch_key

VIEWS = ("group", "teacher", "room")


def entities(placements: list, by: str) -> list:
    if by == "group":
        return sorted({k for p in placements for k in p.all_groups if not is_batch_key(k)})
    if by == "teacher":
        return sorted({p.teacher for p in placements})
    if by == "room":
        return sorted({p.room for p in placements})
    raise ValueError(f"by must be one of {VIEWS}")


def select(placements: list, by: str, name: str) -> list:
    if by == "group":
        return [p for p in placements if name in p.all_groups]
    if by == "teacher":
        return [p for p in placements if p.teacher == name]
    return [p for p in placements if p.room == name]


def grid_matrix(placements: list, cfg: Config, by: str) -> list:
    """Rows = periods, columns = days.  Each cell is a short multi-line text ('' when empty, 'Lunch' for lunch)."""
    cells = defaultdict(list)
    for p in placements:
        who = {"group": f"{p.teacher} · {p.room}", "teacher": f"{p.label or p.group} · {p.room}",
               "room": f"{p.label or p.group} · {p.teacher}"}[by]
        for sl in p.slots():
            first = sl == p.start
            cells[(p.day, sl)].append(f"{'Lab: ' if p.kind == 'lab' and first else ''}{p.subject}\n{who}" if first
                                      else f"(cont.) {p.subject}")
    rows = []
    for sl in range(cfg.num_slots):
        row = [cfg.slot_label(sl)]
        for d in cfg.days:
            if sl in cfg.lunch_slots:
                row.append("Lunch")
            elif not cfg.slot_open(d, sl):
                row.append("-")
            else:
                row.append("\n".join(cells.get((d, sl), [])))
        rows.append(row)
    return rows


def _sheet_name(name: str, used: set) -> str:
    base = "".join(ch for ch in name if ch not in '[]:*?/\\')[:28] or "Sheet"
    n, i = base, 2
    while n.lower() in used:
        n = f"{base[:25]}_{i}"
        i += 1
    used.add(n.lower())
    return n


def to_excel(placements: list, cfg: Config, by: str = "group") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    thin = Side(style="thin", color="BBBBBB")
    for name in entities(placements, by):
        ws = wb.create_sheet(_sheet_name(name, used))
        ws.append([f"{by.capitalize()}: {name}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Period"] + list(cfg.days))
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="334155")
            c.alignment = Alignment(horizontal="center")
        for row in grid_matrix(select(placements, by, name), cfg, by):
            ws.append(row)
        for r in ws.iter_rows(min_row=3):
            for c in r:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if c.value == "Lunch":
                    c.fill = PatternFill("solid", fgColor="E5E7EB")
                elif c.column > 1 and c.value and c.value != "-":
                    c.fill = PatternFill("solid", fgColor="DBEAFE")
        ws.column_dimensions["A"].width = 16
        for i in range(len(cfg.days)):
            ws.column_dimensions[chr(66 + i)].width = 26
        for r in range(3, 3 + cfg.num_slots):
            ws.row_dimensions[r].height = 42
    # a flat list of every session, handy for filtering
    ws = wb.create_sheet("All sessions", 0)
    ws.append(["Attended by", "Day", "Start", "End", "Subject", "Type", "Teacher", "Room", "Building"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for p in sorted(placements, key=lambda p: (cfg.days.index(p.day), p.start, p.label or p.group)):
        ws.append([p.label or p.group, p.day, cfg.fmt(cfg.slot_start(p.start)), cfg.fmt(cfg.slot_end(p.end - 1)),
                   p.subject, p.kind, p.teacher, p.room, p.building])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(placements: list, cfg: Config, by: str = "group", title: str = "Timetable") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak, Spacer
    from xml.sax.saxutils import escape
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
                            title=title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    story = []
    names = entities(placements, by)
    if not names:
        story.append(Paragraph("No sessions to show.", styles["BodyText"]))
    for i, name in enumerate(names):
        story.append(Paragraph(f"{escape(title)} - {by}: {escape(name)}", styles["Heading2"]))
        story.append(Spacer(1, 6))
        data = [["Period"] + list(cfg.days)]
        for row in grid_matrix(select(placements, by, name), cfg, by):
            data.append([Paragraph(escape(c).replace("\n", "<br/>"), cell) for c in row])
        width = (landscape(A4)[0] - 48 - 70) / len(cfg.days)
        t = Table(data, colWidths=[70] + [width] * len(cfg.days), repeatRows=1)
        style = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for r, row in enumerate(grid_matrix(select(placements, by, name), cfg, by), start=1):
            for c, txt in enumerate(row[1:], start=1):
                if txt == "Lunch":
                    style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor("#E5E7EB")))
                elif txt and txt != "-":
                    style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor("#DBEAFE")))
        t.setStyle(TableStyle(style))
        story.append(t)
        if i < len(names) - 1:
            story.append(PageBreak())
    doc.build(story)
    return buf.getvalue()
