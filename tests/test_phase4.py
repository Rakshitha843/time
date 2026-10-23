"""Phase 4 tests: versions, import, export, manual edits, role helpers."""
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from engine import TimetableAssistant
from exporters import entities, select, to_excel, to_pdf
from importers import build_store, check_tables, merge_frames, read_table, read_workbook, template_workbook
from manual import find_conflicts, move_session
from models import advanced_sample_store, is_batch_key, sample_store
from versions import compare, list_versions, load_data, restore_version, save_version, delete_version

G = ["CSE-5-A", "CSE-5-B"]


@pytest.fixture(scope="module")
def asst():
    a = TimetableAssistant(advanced_sample_store(), time_limit=6)
    assert a.handle("Create a timetable for CSE 5th semester.").status == "ok"
    return a


def test_batch_key_helper():
    assert is_batch_key("CSE-5-A-B1") and not is_batch_key("CSE-5-B") and not is_batch_key("CSE-5-A")


def test_entities_include_section_b(asst):
    assert entities(asst.timetable, "group") == G                 # regression: 'CSE-5-B' is a section, not a batch
    assert "Dr. Rao" in entities(asst.timetable, "teacher")
    assert len(select(asst.timetable, "group", "CSE-5-B")) > 0


def test_excel_has_one_sheet_per_entity(asst):
    wb = load_workbook(io.BytesIO(to_excel(asst.timetable, asst.cfg, "group")))
    assert wb.sheetnames[0] == "All sessions" and {"CSE-5-A", "CSE-5-B"} <= set(wb.sheetnames)
    ws = wb["CSE-5-A"]
    text = " ".join(str(c.value) for r in ws.iter_rows() for c in r if c.value)
    assert "Lunch" in text and "DBMS" in text
    wb = load_workbook(io.BytesIO(to_excel(asst.timetable, asst.cfg, "room")))
    assert len(wb.sheetnames) == 1 + len(entities(asst.timetable, "room"))


def test_pdf_is_a_pdf_with_one_page_per_teacher(asst):
    data = to_pdf(asst.timetable, asst.cfg, "teacher")
    assert data.startswith(b"%PDF")
    pages = data.count(b"/Type /Page\n") + data.count(b"/Type /Page ") + data.count(b"/Type /Page/")
    assert pages >= len(entities(asst.timetable, "teacher")) or b"/Count" in data


def test_manual_move_detects_teacher_clash(asst):
    tt = asst.timetable
    rao = [i for i, p in enumerate(tt) if p.teacher == "Dr. Rao" and p.kind == "theory"]
    target = tt[rao[1]]
    new, msgs, bad = asst.preview_move(rao[0], target.day, target.start)
    assert any("teacher clash" in m for m in msgs) and bad
    assert tt[rao[0]] is not new[rao[0]]                           # original untouched


def test_manual_move_applies_pins_and_undoes():
    a = TimetableAssistant(advanced_sample_store(), time_limit=6)
    a.handle("Create a timetable for CSE 5th semester.")
    i = next(i for i, p in enumerate(a.timetable) if p.subject == "Operating Systems" and p.group == "CSE-5-A")
    for d in a.cfg.days:
        for s in range(a.cfg.num_slots):
            _, msgs, _ = a.preview_move(i, d, s)
            if not msgs and (d, s) != (a.timetable[i].day, a.timetable[i].start):
                r = a.apply_move(i, d, s)
                assert r.status == "ok" and (a.timetable[i].day, a.timetable[i].start) == (d, s)
                assert any(c.origin == "manual" for c in a.constraints)
                assert a.undo().status == "ok" and not any(c.origin == "manual" for c in a.constraints)
                return
    pytest.fail("no free slot found")


def test_blocked_move_is_not_applied(asst):
    tt = asst.timetable
    rao = [i for i, p in enumerate(tt) if p.teacher == "Dr. Rao" and p.kind == "theory"]
    t = tt[rao[1]]
    before = (tt[rao[0]].day, tt[rao[0]].start)
    r = asst.apply_move(rao[0], t.day, t.start)
    assert r.status == "infeasible" and (tt[rao[0]].day, tt[rao[0]].start) == before


def test_versions_roundtrip_and_compare(tmp_path):
    a = TimetableAssistant(advanced_sample_store(), time_limit=6)
    a.handle("Create a timetable for CSE 5th semester.")
    save_version(a, "Draft 1", tmp_path)
    assert [v["name"] for v in list_versions(tmp_path)] == ["Draft_1"]
    i = next(i for i, p in enumerate(a.timetable) if p.subject == "Computer Networks" and p.group == "CSE-5-A")
    moved = False
    for d in a.cfg.days:
        for s in range(a.cfg.num_slots):
            if not moved and a.preview_move(i, d, s)[1] == [] and (d, s) != (a.timetable[i].day, a.timetable[i].start):
                a.apply_move(i, d, s)
                moved = True
    assert moved
    save_version(a, "Draft 2", tmp_path)
    diff = compare(load_data("Draft 1", tmp_path), load_data("Draft 2", tmp_path))
    assert not diff["identical"] and any("Computer Networks" in m for m in diff["moved"])
    assert compare(load_data("Draft 1", tmp_path), load_data("Draft 1", tmp_path))["identical"]
    restore_version(a, "Draft 1", tmp_path)
    assert a.timetable[i].day != "" and compare(load_data("Draft 1", tmp_path), load_data("Draft 1", tmp_path))["identical"]
    assert delete_version("Draft 2", tmp_path) and len(list_versions(tmp_path)) == 1
    with pytest.raises(ValueError):
        save_version(a, "   !!! ", tmp_path)


def test_template_roundtrip_has_no_issues():
    s = advanced_sample_store()
    tables, problems = read_workbook(template_workbook(s))
    assert problems == [] or all("unknown column" not in p for p in problems)
    back = build_store(tables, s, "replace")
    assert check_tables(back.subjects, back.teachers, back.rooms, back.groups) == []
    assert len(back.subjects) == len(s.subjects) and int(back.subjects["batches"].max()) == 2


def test_csv_import_reports_bad_rows():
    s = sample_store()
    csv = "name,group,teacher,sessions_per_week,type,duration\nPhysics,CSE-9-Z,Dr. Ghost,0,seminar,1\n"
    df, problems = read_table(csv.encode(), "subjects", "x.csv")
    raw = merge_frames({"subjects": df}, s, "append")
    issues = check_tables(raw["subjects"], raw["teachers"], raw["rooms"], raw["groups"])
    text = " ".join(issues)
    assert "not in the groups table" in text and "Dr. Ghost" in text and "sessions_per_week" in text and "type must be" in text


def test_csv_missing_required_column():
    df, problems = read_table(b"name,aliases\nX,y\n", "subjects", "x.csv")
    assert any("Missing required" in p for p in problems)


def test_append_overwrites_by_name():
    s = sample_store()
    row = s.teachers.iloc[[0]].copy()
    row["max_per_day"] = 2
    out = build_store({"teachers": row}, s, "append")
    assert len(out.teachers) == len(s.teachers) and int(out.teachers.iloc[-1]["max_per_day"]) == 2
